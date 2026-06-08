import os
import re
import requests
import xlrd
from flask import Flask, request, jsonify
from flask_cors import CORS
from openpyxl import load_workbook
from werkzeug.utils import secure_filename

app = Flask(__name__)
CORS(app, origins="*")

BACKEND_URL = os.getenv("BACKEND_URL", "https://ugel-1bb.onrender.com")
UPLOAD_FOLDER = "/app/uploads"
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

DNI_PATTERN = re.compile(r"DNI\s*(\d+)", re.IGNORECASE)
IGNORAR_CONCEPTOS = {"REINTEGRO", "ESCOLARIDAD", "AGUINALDO", "DETALLE"}


def parsear_valor(v):
    if v is None:
        return None
    try:
        f = float(v)
        return round(f, 2) if f else None
    except (TypeError, ValueError):
        return None


def limpiar_concepto(c: str) -> str:
    return c.strip().lstrip("+").strip().upper()


def analizar_duplicados(empleados: list) -> dict:
    dni_count = {}
    nombre_count = {}
    dni_empleados = {}
    nombre_empleados = {}
    
    for emp in empleados:
        dni = emp.get("dni")
        nombre = emp.get("nombre")
        monto = emp.get("total_liquido")
        
        if dni:
            if dni not in dni_count:
                dni_count[dni] = 0
                dni_empleados[dni] = []
            dni_count[dni] += 1
            dni_empleados[dni].append({"nombre": nombre, "monto": monto})
        
        if nombre:
            if nombre not in nombre_count:
                nombre_count[nombre] = 0
                nombre_empleados[nombre] = []
            nombre_count[nombre] += 1
            nombre_empleados[nombre].append({"dni": dni, "monto": monto})
    
    dnis_duplicados = [(dni, count) for dni, count in dni_count.items() if count > 1]
    nombres_duplicados = [(nombre, count) for nombre, count in nombre_count.items() if count > 1]
    
    return {
        "dnis_duplicados": [{"dni": dni, "count": count, "empleados": dni_empleados[dni]} for dni, count in dnis_duplicados],
        "nombres_duplicados": [{"nombre": nombre, "count": count, "empleados": nombre_empleados[nombre]} for nombre, count in nombres_duplicados],
    }


def calcular_monto_total(empleados: list) -> float:
    total = 0.0
    for emp in empleados:
        liquido = emp.get("total_liquido")
        if liquido:
            total += float(liquido)
    return round(total, 2)


def extraer_dni(texto: str):
    if not texto:
        return None
    m = DNI_PATTERN.search(str(texto))
    return m.group(1) if m else None


def _leer_filas(filepath: str) -> list:
    """Return all rows as a list of tuples, supporting both .xls and .xlsx."""
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".xls":
        wb = xlrd.open_workbook(filepath)
        ws = wb.sheet_by_index(0)
        filas = []
        for rx in range(ws.nrows):
            row = []
            for cx in range(ws.ncols):
                cell = ws.cell(rx, cx)
                # xlrd type 0=empty,1=text,2=number,3=date,4=bool,5=error,6=blank
                if cell.ctype in (0, 5, 6):
                    row.append(None)
                elif cell.ctype == 2:
                    # Return int if value is whole number, else float
                    v = cell.value
                    row.append(int(v) if v == int(v) else v)
                else:
                    row.append(cell.value)
            filas.append(tuple(row))
        return filas
    else:
        wb = load_workbook(filepath, data_only=True)
        ws = wb.active
        return list(ws.iter_rows(values_only=True))


def extraer_empleados(filepath: str) -> list:
    """
    Parse a payroll Excel with vertical-block layout:

      Col A    | Col B           | Col C    | Col D
      ---------+-----------------+----------+--------
      COLEGIO  | (school name)   | DETALLE  | MES
      HABERES  | Apellidos Nomb. | BASICA   | 0.03
      (merged) | (merged)        | DL19990  | 60.00
               | CARGO/PUESTO    | TPH      | 19.20
               | RD 150-93       | ...
               | uu-01-0-005     |
               | DISTRITO: xxx   |          |
      DSCTOS   |                 | DL20530  | 3.80
      TOTAL HABERES              |          | 153.92
      TOTAL DESCUENTOS           |          | 76.27
      TOTAL LIQUIDO              |          | 77.65

    Mes/Año are NOT read from the file; they come from the user request.
    """
    filas = _leer_filas(filepath)

    empleados = []
    i = 0
    n = len(filas)
    colegio_actual = None

    while i < n:
        fila = filas[i]
        col_a = str(fila[0]).strip() if fila[0] else ""

        # Colegio: row before HABERES has school name in Col B
        col_b = str(fila[1]).strip() if len(fila) > 1 and fila[1] else ""
        if col_b and i + 1 < n:
            siguiente = str(filas[i + 1][0]).strip() if filas[i + 1][0] else ""
            if siguiente == "HABERES":
                colegio_actual = col_b

        if col_a == "HABERES" and len(fila) > 1 and fila[1]:
            emp = {
                "nombre": str(fila[1]).strip(),
                "cargo": None,
                "resolucion": None,
                "codigo": None,
                "dni": None,
                "colegio": colegio_actual,
                "distrito": None,
                "haberes": [],
                "descuentos": [],
                "total_haberes": None,
                "total_descuentos": None,
                "total_liquido": None,
            }
            colegio_actual = None

            concepto_inicial = str(fila[2]).strip() if len(fila) > 2 and fila[2] else ""
            valor_inicial = parsear_valor(fila[3] if len(fila) > 3 else None)
            if concepto_inicial and limpiar_concepto(concepto_inicial) not in IGNORAR_CONCEPTOS:
                if valor_inicial is not None:
                    emp["haberes"].append({"concepto": concepto_inicial, "monto": valor_inicial})

            seccion = "HABERES"
            i += 1

            while i < n:
                f = filas[i]
                a = str(f[0]).strip() if len(f) > 0 and f[0] else ""
                b = str(f[1]).strip() if len(f) > 1 and f[1] else ""
                c_cell = str(f[2]).strip() if len(f) > 2 and f[2] else ""
                d = parsear_valor(f[3] if len(f) > 3 else None)

                if a == "TOTAL HABERES":
                    emp["total_haberes"] = d
                    i += 1
                    continue
                if a == "TOTAL DESCUENTOS":
                    emp["total_descuentos"] = d
                    i += 1
                    continue
                if a == "TOTAL LIQUIDO":
                    emp["total_liquido"] = d
                    i += 1
                    break

                if a == "DSCTOS":
                    seccion = "DSCTOS"
                    if c_cell and limpiar_concepto(c_cell) not in IGNORAR_CONCEPTOS and d is not None:
                        emp["descuentos"].append({"concepto": c_cell, "monto": d})
                    i += 1
                    continue

                if a == "HABERES" and len(f) > 1 and f[1]:
                    break

                if b:
                    dni_encontrado = extraer_dni(b)
                    if dni_encontrado:
                        emp["dni"] = dni_encontrado
                    elif re.match(r"^(DISTRITO|DIST\.?)\s*:?\s*", b, re.IGNORECASE):
                        dist = re.sub(r"^(DISTRITO|DIST\.?)\s*:?\s*", "", b, flags=re.IGNORECASE).strip()
                        if dist:
                            emp["distrito"] = dist
                    elif re.match(r"^(RD|RM|DS|LEY|DL|R\.D\.|R\.M\.)\w*\s*[\d\-./A-Za-z]+", b, re.IGNORECASE):
                        emp["resolucion"] = b
                    elif re.match(r"^uu-", b, re.IGNORECASE):
                        emp["codigo"] = b
                    elif not re.match(r"^(TOTAL|DSCTOS)", b, re.IGNORECASE):
                        # Order: distrito first, then cargo
                        if emp["distrito"] is None:
                            emp["distrito"] = b
                        elif emp["cargo"] is None:
                            emp["cargo"] = b

                if c_cell:
                    nombre_limpio = limpiar_concepto(c_cell)
                    if nombre_limpio not in IGNORAR_CONCEPTOS and not nombre_limpio.startswith("+"):
                        if d is not None:
                            item = {"concepto": c_cell.strip(), "monto": d}
                            if seccion == "HABERES":
                                emp["haberes"].append(item)
                            else:
                                emp["descuentos"].append(item)

                i += 1

            empleados.append(emp)
            continue

        i += 1

    return empleados


# ─────────────────────────────────────────────────────────────────────────────

@app.route("/health", methods=["GET"])
def health():
    return jsonify({"status": "ok"})


@app.route("/process-excel", methods=["POST"])
def process_excel():
    if "file" not in request.files:
        return jsonify({"error": "No se encontró archivo"}), 400

    file = request.files["file"]
    if not file.filename:
        return jsonify({"error": "Archivo sin nombre"}), 400

    filename = secure_filename(file.filename)
    if not filename.lower().endswith(('.xlsx', '.xls')):
        return jsonify({"error": "Solo se aceptan archivos Excel (.xlsx, .xls)"}), 400

    mes = request.form.get("mes", type=int)
    anio = request.form.get("anio", type=int)

    if not mes or not anio:
        return jsonify({"error": "Se requieren los campos mes y anio"}), 400
    if not (1 <= mes <= 12):
        return jsonify({"error": "Mes debe estar entre 1 y 12"}), 400

    filepath = os.path.join(app.config["UPLOAD_FOLDER"], filename)
    file.save(filepath)

    try:
        empleados = extraer_empleados(filepath)
        
        analisis = analizar_duplicados(empleados)
        monto_total = calcular_monto_total(empleados)

        payload = {
            "mes": mes,
            "anio": anio,
            "total_empleados": len(empleados),
            "empleados": empleados,
        }

        response = requests.post(
            f"{BACKEND_URL}/api/importar/haberes",
            json=payload,
            timeout=300,
        )

        if response.status_code == 200:
            resp_data = response.json()
            return jsonify({
                "message": "Excel procesado correctamente",
                "personal_creados": resp_data.get("personal_creados", 0),
                "planillas_creadas": resp_data.get("planillas_creadas", 0),
                "personal": len(empleados),
                "planillas": len(empleados),
                "errores": resp_data.get("errores", []),
                "personal_actualizados": resp_data.get("personal_actualizados", 0),
                "dnis_duplicados": analisis["dnis_duplicados"],
                "nombres_duplicados": analisis["nombres_duplicados"],
                "monto_total": monto_total,
                "duplicados": resp_data.get("duplicados", []),
            })
        else:
            return jsonify({
                "error": "Error al enviar al backend",
                "details": response.text,
            }), response.status_code

    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        if os.path.exists(filepath):
            os.remove(filepath)


@app.route("/validate-excel", methods=["POST"])
def validate_excel():
    if "file" not in request.files:
        return jsonify({"error": "No se encontró archivo"}), 400

    file = request.files["file"]
    if not file.filename:
        return jsonify({"error": "Archivo sin nombre"}), 400

    filename = secure_filename(file.filename)
    if not filename.lower().endswith(('.xlsx', '.xls')):
        return jsonify({"error": "Solo se aceptan archivos Excel (.xlsx, .xls)"}), 400

    filepath = os.path.join(app.config["UPLOAD_FOLDER"], filename)
    file.save(filepath)

    try:
        empleados = extraer_empleados(filepath)
        analisis = analizar_duplicados(empleados)
        monto_total = calcular_monto_total(empleados)
        
        return jsonify({
            "valid": True,
            "total_empleados": len(empleados),
            "preview": empleados[:5],
            "dnis_duplicados": analisis["dnis_duplicados"],
            "nombres_duplicados": analisis["nombres_duplicados"],
            "monto_total": monto_total,
        })
    except Exception as e:
        return jsonify({"valid": False, "error": str(e)}), 400
    finally:
        if os.path.exists(filepath):
            os.remove(filepath)


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=8081)
